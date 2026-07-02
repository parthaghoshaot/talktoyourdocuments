import os
import fitz
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".doc", ".pptx", ".ppt", ".xlsx", ".xls"}
MIN_IMAGE_BYTES = 5000
MAX_IMAGES_PER_DOC = 20


def parse_document(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    parsers = {
        ".pdf": parse_pdf, ".docx": parse_docx, ".doc": parse_docx,
        ".pptx": parse_pptx, ".ppt": parse_pptx, ".xlsx": parse_xlsx, ".xls": parse_xlsx,
    }
    parser = parsers.get(ext)
    if not parser:
        raise ValueError(f"Unsupported file type: {ext}")
    return parser(file_path)


def parse_pdf(file_path):
    doc = fitz.open(file_path)
    text_parts = []
    images = []
    for page_num, page in enumerate(doc, 1):
        text_parts.append(f"[Page {page_num}]")
        blocks = page.get_text("dict", flags=fitz.TEXT_PRESERVE_WHITESPACE)["blocks"]
        page_lines = []
        for block in blocks:
            if block["type"] != 0:
                continue
            for line in block.get("lines", []):
                spans = line.get("spans", [])
                if not spans:
                    continue
                text = "".join(s["text"] for s in spans).strip()
                if not text:
                    continue
                max_size = max(s["size"] for s in spans)
                is_bold = any("bold" in s.get("font", "").lower() for s in spans)
                if max_size >= 14 and len(text) < 120:
                    page_lines.append(f"[H1] {text}")
                elif (max_size >= 12 or is_bold) and len(text) < 120:
                    page_lines.append(f"[H2] {text}")
                else:
                    page_lines.append(text)
        text_parts.append("\n".join(page_lines))
        if len(images) < MAX_IMAGES_PER_DOC:
            for img_info in page.get_images(full=True):
                try:
                    pix = fitz.Pixmap(doc, img_info[0])
                    if pix.n - pix.alpha > 3:
                        pix = fitz.Pixmap(fitz.csRGB, pix)
                    png = pix.tobytes("png")
                    if len(png) >= MIN_IMAGE_BYTES:
                        images.append({"data": png, "page": page_num, "mime": "image/png"})
                except Exception:
                    continue
    doc.close()
    return "\n".join(text_parts), images


def parse_docx(file_path):
    doc = Document(file_path)
    text_parts = []
    images = []
    for para in doc.paragraphs:
        if not para.text.strip():
            continue
        style = (para.style.name or "").lower()
        if "heading 1" in style or "title" in style:
            text_parts.append(f"[H1] {para.text}")
        elif "heading 2" in style:
            text_parts.append(f"[H2] {para.text}")
        elif "heading" in style:
            text_parts.append(f"[H3] {para.text}")
        else:
            text_parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                text_parts.append(row_text)
    for rel in doc.part.rels.values():
        if "image" in rel.reltype and len(images) < MAX_IMAGES_PER_DOC:
            try:
                blob = rel.target_part.blob
                if len(blob) >= MIN_IMAGE_BYTES:
                    images.append({"data": blob, "page": None, "mime": "image/png"})
            except Exception:
                continue
    return "\n".join(text_parts), images


def parse_pptx(file_path):
    prs = Presentation(file_path)
    text_parts = []
    images = []
    for slide_num, slide in enumerate(prs.slides, 1):
        slide_texts = [f"[Slide {slide_num}]"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    if para.text.strip():
                        slide_texts.append(para.text)
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        slide_texts.append(row_text)
            if shape.shape_type == 13 and len(images) < MAX_IMAGES_PER_DOC:
                try:
                    blob = shape.image.blob
                    if len(blob) >= MIN_IMAGE_BYTES:
                        images.append({"data": blob, "page": slide_num, "mime": "image/png"})
                except Exception:
                    continue
        text_parts.append("\n".join(slide_texts))
    return "\n\n".join(text_parts), images


def parse_xlsx(file_path):
    wb = load_workbook(file_path, read_only=True, data_only=True)
    text_parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        text_parts.append(f"[Sheet: {sheet_name}]")
        for row in ws.iter_rows(values_only=True):
            row_text = " | ".join(str(cell) for cell in row if cell is not None)
            if row_text.strip():
                text_parts.append(row_text)
    wb.close()
    return "\n".join(text_parts), []


def get_supported_files(directory):
    files = []
    for root, _, filenames in os.walk(directory):
        for fname in filenames:
            if os.path.splitext(fname)[1].lower() in SUPPORTED_EXTENSIONS:
                files.append(os.path.join(root, fname))
    return files
